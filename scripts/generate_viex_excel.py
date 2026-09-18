#!/usr/bin/env python3
"""Create the validated two-sheet VIEX workbook from JSON records."""
import json
import sys
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROSPECTOS = ["ID Prospecto","Tipo","Categoría","Nombre de fantasía","Razón social","RUT","Giro","Actividad","Descripción","Servicios","Público objetivo","Región","Ciudad","Comuna","Dirección","Horario","Sitio web","Google Business","LinkedIn empresa","Contacto","Cargo","Teléfono","WhatsApp","Email","LinkedIn contacto","Otras redes","Oportunidad web","Oportunidad SEO","Oportunidad UX/CRO","Oportunidad Ads","Oportunidad automatización","Oportunidad analítica","Oportunidad comercial","Motivo de contacto","Score","Prioridad","Estado","Fecha investigación","Fuentes"]
LISTA = ["ID Prospecto","Tipo","Categoría","Nombre de fantasía / Persona","Razón social","RUT","Giro","Actividad","Región","Ciudad","Comuna","Dirección","Sitio web","Google Business","Cantidad de contactos","Score","Prioridad","Oportunidad comercial","Motivo de contacto","Estado","Fecha investigación"]
VALID_PRIORIDADES = {"A", "B", "C", "D"}
VALID_ESTADOS = {"Nuevo", "Investigado", "Priorizado", "Por contactar", "Contactado", "En conversación", "Oportunidad", "Cliente", "No interesado", "No contactar", "Información insuficiente"}
URL_FIELDS = {"Sitio web", "Google Business", "LinkedIn empresa", "LinkedIn contacto", "Otras redes"}


def val(row, key):
    value = row.get(key)
    return "No encontrado" if value is None or value == "" else value


def url(value):
    parsed = urlparse(str(value))
    return str(value) if parsed.scheme in {"http", "https"} and parsed.netloc else None


def validate(contacts, master):
    ids = [str(val(row, "ID Prospecto")) for row in master]
    if len(ids) != len(set(ids)):
        raise ValueError("IDs duplicados en LISTA_PROSPECTOS")
    master_ids = set(ids)
    counts = {}
    for row in contacts:
        prospect_id = str(val(row, "ID Prospecto"))
        if prospect_id not in master_ids:
            raise ValueError(f"ID de contacto ausente en lista maestra: {prospect_id}")
        counts[prospect_id] = counts.get(prospect_id, 0) + 1
        score = val(row, "Score")
        if not isinstance(score, int) or not 0 <= score <= 100:
            raise ValueError(f"Score inválido para {prospect_id}: {score!r}")
        if val(row, "Prioridad") not in VALID_PRIORIDADES:
            raise ValueError(f"Prioridad inválida para {prospect_id}")
        if val(row, "Estado") not in VALID_ESTADOS:
            raise ValueError(f"Estado inválido para {prospect_id}")
    for row in master:
        prospect_id = str(val(row, "ID Prospecto"))
        if val(row, "Cantidad de contactos") != counts.get(prospect_id, 0):
            raise ValueError(f"Cantidad de contactos inconsistente para {prospect_id}")


def make_rows(records, headers):
    return [[val(row, header) for header in headers] for row in records]


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: generate_viex_excel.py records.json output.xlsx")
    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
    contacts = data.get("prospectos", data) if isinstance(data, dict) else data
    master = data.get("lista_prospectos") if isinstance(data, dict) else None
    if not isinstance(contacts, list) or not contacts:
        raise ValueError("El JSON debe contener una lista no vacía de prospectos")
    if master is None:
        grouped = {}
        for row in contacts:
            grouped.setdefault(str(val(row, "ID Prospecto")), []).append(row)
        master = []
        for rows in grouped.values():
            row = dict(rows[0])
            row["Cantidad de contactos"] = len(rows)
            row["Nombre de fantasía / Persona"] = val(row, "Nombre de fantasía")
            master.append(row)
    validate(contacts, master)

    workbook = Workbook()
    sheets = [(workbook.active, "PROSPECTOS", PROSPECTOS, contacts), (workbook.create_sheet(), "LISTA_PROSPECTOS", LISTA, master)]
    for sheet, title, headers, records in sheets:
        sheet.title = title
        sheet.append(headers)
        for row in make_rows(records, headers):
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            width = min(48, max(len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 80) + 1)) + 2)
            sheet.column_dimensions[letter].width = max(12, width)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if headers[cell.column - 1] in URL_FIELDS:
                    link = url(cell.value)
                    if link:
                        cell.hyperlink = link
                        cell.style = "Hyperlink"
        table_ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        table = Table(displayName=f"T_{title}", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    workbook.save(sys.argv[2])


if __name__ == "__main__":
    main()
